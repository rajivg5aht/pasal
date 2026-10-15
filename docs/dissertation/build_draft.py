"""Build an editable evidence-based draft; never fabricate final-result sections."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.8)
section.left_margin = section.right_margin = Inches(.9)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Calibri', Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for name, size in [('Title',24),('Heading 1',18),('Heading 2',13)]:
    style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.keep_with_next=True
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Pasal working draft | ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.add_paragraph('Generative Artificial Intelligence Enhanced Hybrid Recommender System for Cold Start Mitigation in Online Retail','Title')
doc.add_paragraph('Pasal thesis working draft','Subtitle')
doc.add_paragraph('Rajiv Ganeju | Student ID 240009')
doc.add_paragraph('Prepared 15 October 2026')
doc.add_paragraph('Status: not submission-ready. This editable draft assembles verified methodology, implementation and development evidence. Final holdout results, the completed generation sample and final conclusions remain outstanding. No abstract claiming findings is supplied.')
doc.add_paragraph('The approved title retains the hyphenation “Generative Artificial Intelligence-Enhanced Hybrid Recommender System for Cold-Start Mitigation in Online Retail.” The cover typography above uses the same wording without punctuation.')
doc.add_paragraph('Contents','Heading 1')
for text in ['1 Introduction','2 Related work','3 Methodology','4 System implementation','5 Development results and discussion','Verified references']:
    doc.add_paragraph(text)
doc.add_paragraph('Institution-specific word limits, citation style and submission template have not been supplied. Final formatting must be checked against those requirements.')

def clean(text):
    return re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'\1 (\2)',text).replace('**','').replace('`','')

def add_markdown(path):
    lines=path.read_text(encoding='utf-8').splitlines(); paragraph=[];table=[]
    def flush():
        if paragraph:doc.add_paragraph(clean(' '.join(paragraph)));paragraph.clear()
    def flush_table():
        if not table:return
        rows=[r for r in table if not all(re.fullmatch(r':?-+:?',c) for c in r)]
        t=doc.add_table(rows=1, cols=len(rows[0]));t.style='Light Shading Accent 1'
        for cell,value in zip(t.rows[0].cells,rows[0]):cell.text=clean(value)
        for row in rows[1:]:
            for cell,value in zip(t.add_row().cells,row):cell.text=clean(value)
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:run.font.size=Pt(9)
        table.clear()
    for line in lines+['']:
        if line.startswith('|'):
            flush();table.append([s.strip() for s in line.strip('|').split('|')]);continue
        flush_table()
        if not line.strip():flush();continue
        if line.startswith('#'):
            flush();level=len(line)-len(line.lstrip('#'));doc.add_paragraph(clean(line.lstrip('#').strip()),f'Heading {min(level,2)}')
        elif line.startswith('- '):
            flush();doc.add_paragraph(clean(line[2:]),'List Bullet')
        else:paragraph.append(line.strip())

for name in ['01-introduction.md','02-literature.md','03-methodology.md','04-implementation.md','05-results-discussion.md','references.md']:
    doc.add_page_break();add_markdown(ROOT/name)
doc.core_properties.title='Pasal thesis working draft'
doc.core_properties.author='Rajiv Ganeju'
out=ROOT/'Pasal_Thesis_Working_Draft.docx'
doc.save(out)
print(out)
